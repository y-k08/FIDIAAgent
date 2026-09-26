#!/usr/bin/env python3
"""文部科学省「令和7年度全国大学一覧」Excelから universities.json を生成する。"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT / "data" / "universities.json"
OUTPUT_JS_PATH = ROOT / "js" / "universities-data.js"

TYPE_NATIONAL = "国立"
TYPE_PUBLIC = "公立"
TYPE_PRIVATE = "私立"
TYPE_BROADCAST = "放送大学"
TYPE_ORDER = (TYPE_NATIONAL, TYPE_PUBLIC, TYPE_PRIVATE, TYPE_BROADCAST)

PREFECTURES = (
    "北海道",
    "青森県",
    "岩手県",
    "宮城県",
    "秋田県",
    "山形県",
    "福島県",
    "茨城県",
    "栃木県",
    "群馬県",
    "埼玉県",
    "千葉県",
    "東京都",
    "神奈川県",
    "新潟県",
    "富山県",
    "石川県",
    "福井県",
    "山梨県",
    "長野県",
    "岐阜県",
    "静岡県",
    "愛知県",
    "三重県",
    "滋賀県",
    "京都府",
    "大阪府",
    "兵庫県",
    "奈良県",
    "和歌山県",
    "鳥取県",
    "島根県",
    "岡山県",
    "広島県",
    "山口県",
    "徳島県",
    "香川県",
    "愛媛県",
    "高知県",
    "福岡県",
    "佐賀県",
    "長崎県",
    "熊本県",
    "大分県",
    "宮崎県",
    "鹿児島県",
    "沖縄県",
)


def find_raw_dir() -> Path:
    for candidate in (ROOT / "data" / "raw", ROOT / "data:raw:"):
        if candidate.is_dir() and any(candidate.glob("*.xlsx")):
            return candidate
    raise SystemExit("Excelが見つかりません。data/raw/ にxlsxを置いてください。")


def classify_file(path: Path, sheet_names: list[str]) -> str | None:
    name = path.name
    if "_04." in name or "索引" in sheet_names:
        return "index"
    if "_05." in name or "放送大学" in sheet_names:
        return "broadcast"
    if "_01." in name:
        return "national"
    if "_02." in name:
        return "public"
    return None


def first_nonempty(row) -> str:
    for cell in row:
        if cell not in (None, ""):
            return str(cell)
    return ""


def clean_name(heading: str) -> str:
    text = heading.strip()
    text = re.sub(r"（[^）]*）", "", text)
    text = re.sub(r"\([^)]*\)", "", text)
    text = re.sub(r"^(国立|公立|私立)[\s\u3000]*", "", text.strip())
    return text.replace(" ", "").replace("\u3000", "")


def extract_pref(cells) -> str:
    for cell in cells:
        if cell in (None, ""):
            continue
        text = str(cell).replace(" ", "").replace("\u3000", "")
        for pref in PREFECTURES:
            if text.startswith(pref):
                return pref
    return ""


def make_record(name: str, pref: str, kind: str) -> dict[str, str] | None:
    name = name.replace(" ", "").replace("\u3000", "")
    if not name:
        return None
    return {"name": name, "kana": "", "pref": pref, "type": kind}


def read_preview_rows(ws, max_row: int = 30) -> list[tuple]:
    rows = []
    for i, row in enumerate(ws.iter_rows(max_row=max_row, values_only=True), start=1):
        rows.append(row)
        if i >= max_row:
            break
    return rows


def parse_detail_sheet(ws, kind: str) -> dict[str, str] | None:
    rows = read_preview_rows(ws)
    if not rows:
        return None
    heading = first_nonempty(rows[0])
    if not heading:
        return None
    name = clean_name(heading)
    cells = [cell for row in rows for cell in row]
    pref = extract_pref(cells)
    return make_record(name, pref, kind)


def parse_index_sheet(ws) -> list[dict[str, str]]:
    rows = list(ws.iter_rows(values_only=True))
    header_row = None
    header_index = None
    for i, row in enumerate(rows):
        values = [str(cell).strip() if cell not in (None, "") else "" for cell in row]
        if "学校名称" in values and "設置区分" in values:
            header_row = values
            header_index = i
            break
    if header_row is None:
        nonempty = [cell for cell in rows[0] if cell not in (None, "")] if rows else []
        raise SystemExit(f"索引シートの列名が想定と違います。先頭行: {nonempty}")

    name_col = header_row.index("学校名称")
    type_col = header_row.index("設置区分")
    records = []
    for row in rows[header_index + 1 :]:
        values = list(row)
        name = values[name_col] if len(values) > name_col else None
        kind = values[type_col] if len(values) > type_col else None
        if kind not in (None, "") and kind != TYPE_PRIVATE:
            raise SystemExit(f"索引の設置区分が想定外です: {kind!r}")
        record = make_record(str(name) if name not in (None, "") else "", "", TYPE_PRIVATE)
        if record:
            records.append(record)
    return records


def add_unique(records: list[dict[str, str]], seen: set[str], record: dict[str, str]) -> None:
    if record["name"] in seen:
        return
    seen.add(record["name"])
    records.append(record)


def main() -> None:
    raw_dir = find_raw_dir()
    files = sorted(raw_dir.glob("*.xlsx"))
    if not files:
        raise SystemExit(f"{raw_dir} にxlsxがありません。")

    classified: dict[str, Path] = {}
    for path in files:
        wb = load_workbook(path, read_only=True, data_only=True)
        kind = classify_file(path, wb.sheetnames)
        wb.close()
        if kind is None:
            continue
        classified[kind] = path

    records: list[dict[str, str]] = []
    seen: set[str] = set()

    for kind in ("national", "public", "index", "broadcast"):
        path = classified.get(kind)
        if path is None:
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        if kind == "index":
            if "索引" not in wb.sheetnames:
                wb.close()
                raise SystemExit(f"{path.name} に「索引」シートがありません。シート: {wb.sheetnames}")
            for record in parse_index_sheet(wb["索引"]):
                add_unique(records, seen, record)
        elif kind == "broadcast":
            sheet_name = "放送大学" if "放送大学" in wb.sheetnames else wb.sheetnames[0]
            record = parse_detail_sheet(wb[sheet_name], TYPE_BROADCAST)
            if record:
                add_unique(records, seen, record)
        else:
            univ_type = TYPE_NATIONAL if kind == "national" else TYPE_PUBLIC
            for sheet_name in wb.sheetnames:
                record = parse_detail_sheet(wb[sheet_name], univ_type)
                if record:
                    add_unique(records, seen, record)
        wb.close()

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    OUTPUT_JS_PATH.write_text(
        "window.FIDIA_UNIVERSITIES = " + json.dumps(records, ensure_ascii=False) + ";\n",
        encoding="utf-8",
    )

    counts = {kind: 0 for kind in TYPE_ORDER}
    for record in records:
        counts[record["type"]] = counts.get(record["type"], 0) + 1

    print(f"出力: {OUTPUT_PATH.relative_to(ROOT)}")
    print(f"出力: {OUTPUT_JS_PATH.relative_to(ROOT)}")
    print(f"合計: {len(records)}件")
    for kind in TYPE_ORDER:
        print(f"  {kind}: {counts.get(kind, 0)}件")
    extra = set(counts) - set(TYPE_ORDER)
    for kind in sorted(extra):
        print(f"  {kind}: {counts[kind]}件")
    print("先頭10件:")
    for record in records[:10]:
        print(f"  {record}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"エラー: {exc}", file=sys.stderr)
        raise
